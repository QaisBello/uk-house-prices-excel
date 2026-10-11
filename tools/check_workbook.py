"""Recalculate a copy of the workbook in LibreOffice and check its numbers against an independent pandas calculation.

    python tools/check_workbook.py      (needs LibreOffice: soffice on PATH)

Writes data/clean/*.csv (the calculated tables, for the web page) and docs/check_results.md.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RECALC = os.path.join(ROOT, "tools", "recalc_lo.py")
tmp = os.path.join(tempfile.mkdtemp(), "check.xlsx")
shutil.copy(os.path.join(ROOT, "UK_House_Prices.xlsx"), tmp)
res = json.loads(subprocess.run([sys.executable, RECALC, tmp, "120"], capture_output=True, text=True).stdout)
assert res.get("status") == "success", res
wb = load_workbook(tmp, data_only=True)


def table(ws, header_row, ncols):
    rows = list(ws.iter_rows(min_row=header_row, max_col=ncols, values_only=True))
    df = pd.DataFrame(rows[1:], columns=rows[0])
    return df[df.iloc[:, 0].notna()]


q = table(wb["Quarterly"], 4, 14)
an = table(wb["Annual"], 4, 7)
de = table(wb["Decades"], 4, 10)
ck = table(wb["Checks"], 4, 4)

# ---- independent calculation
hp = pd.read_csv(os.path.join(ROOT, "data/raw/nationwide_uk_house_prices.csv"), parse_dates=["Date"])
cpi = pd.read_csv(os.path.join(ROOT, "data/raw/worldbank_cpi.csv"))
uk = cpi[cpi["Country Code"] == "GBR"].set_index("Year")["CPI"] / 100
idx = (1 + uk).cumprod() * 100
idx = idx / idx[idx.index.max()] * 100                      # base year = latest (the workbook default)
hp["Year"] = hp.Date.dt.year
hp["real"] = hp["Price (All)"] * 100 / hp.Year.map(idx)
annual = hp.dropna(subset=["real"]).groupby("Year").agg(price=("Price (All)", "mean"), real=("real", "mean"))
hp["peak"] = hp.real.cummax()
hp["fall"] = hp.real / hp.peak - 1

results = []


def check(label, got, want, tol):
    ok = bool(np.allclose(np.asarray(got, float), np.asarray(want, float), atol=tol, equal_nan=True))
    results.append((label, ok))
    print(("PASS " if ok else "FAIL ") + label)


qv = q.set_index(q.columns[0])
check("Real price, every quarter (to £0.01)", pd.to_numeric(q.iloc[:, 7], errors="coerce").values, hp.real.values, 0.01)
check("Fall from peak, every quarter", pd.to_numeric(q.iloc[:, 13], errors="coerce").values, hp.fall.values, 1e-9)
check("Annual cash price", an.iloc[:, 1].astype(float).values, annual.price.values, 0.01)
check("Annual real price", an.iloc[:, 2].astype(float).values, annual.real.values, 0.01)
check("All Checks sheet tests PASS", [s == "PASS" for s in ck.iloc[:, 2]], [True] * len(ck), 0)

mo = wb["Mortgage"]
loan = mo["D12"].value
r, n = mo["D7"].value / 12, mo["D9"].value * 12
check("Monthly repayment = annuity formula", mo["D14"].value, loan * r / (1 - (1 + r) ** -n), 0.005)

# ---- clean exports for the web page
os.makedirs(os.path.join(ROOT, "data/clean"), exist_ok=True)
qe = q.copy()
qe.columns = ["date", "year", "quarter", "label", "price", "price_new", "cpi_index", "real_price", "yoy", "real_yoy",
              "published_yoy", "diff_pp", "real_peak", "real_fall"]
qe["date"] = pd.to_datetime(qe.date).dt.strftime("%Y-%m-%d")
qe.replace("", np.nan).to_csv(os.path.join(ROOT, "data/clean/quarterly.csv"), index=False, float_format="%.6g")
an.columns = ["year", "price", "real_price", "inflation", "growth", "real_growth", "real_vs_1960"]
an.to_csv(os.path.join(ROOT, "data/clean/annual.csv"), index=False, float_format="%.6g")
de.columns = ["decade", "from", "to", "price_start", "price_end", "growth", "real_growth", "growth_pa", "real_growth_pa", "inflation_pa"]
de.to_csv(os.path.join(ROOT, "data/clean/decades.csv"), index=False, float_format="%.6g")
db = wb["Dashboard"]
falls = [{"episode": db[f"B{r}"].value, "fall": db[f"C{r}"].value, "trough": db[f"E{r}"].value} for r in (17, 18, 19)]
summary = {"latest_price": db["B8"].value, "latest_label": db["B9"].value, "cash_multiple": db["E8"].value, "real_multiple": db["H8"].value,
           "real_growth_pa": db["B12"].value, "inflation_pa": db["E12"].value, "worst_fall": db["H12"].value, "worst_trough": db["H13"].value,
           "falls": falls, "checks": db["F5"].value, "formulas": res["total_formulas"], "base_year": db["C5"].value,
           "mortgage": {k: [mo[f"C{r_}"].value, mo[f"D{r_}"].value] for k, r_ in (("quarter", 6), ("price", 10), ("payment", 14), ("payment_real", 16))}}
json.dump(summary, open(os.path.join(ROOT, "data/clean/summary.json"), "w"), indent=1, default=str)
ck.columns = ["check", "result", "status", "proves"]
ck.to_csv(os.path.join(ROOT, "data/clean/checks.csv"), index=False)

lines = ["# Check results", "", f"Workbook recalculated in LibreOffice: {res['total_formulas']} formulas, {res['total_errors']} errors.", "",
         "Independent recalculation in pandas (`tools/check_workbook.py`):", ""] + [f"- {'PASS' if ok else 'FAIL'}: {l}" for l, ok in results]
open(os.path.join(ROOT, "docs/check_results.md"), "w").write("\n".join(lines) + "\n")
print(json.dumps(summary, indent=1, default=str))
sys.exit(0 if all(ok for _, ok in results) else 1)
