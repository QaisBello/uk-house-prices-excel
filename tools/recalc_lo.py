"""Recalculate an .xlsx with LibreOffice (headless) and count formula errors.

    python tools/recalc_lo.py book.xlsx      # rewrites book.xlsx with calculated values, prints JSON
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

ERRORS = ("#NAME?", "#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")
path = os.path.abspath(sys.argv[1])
out = tempfile.mkdtemp()
profile = tempfile.mkdtemp()
subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                "--outdir", out, path], check=True, capture_output=True, timeout=int(sys.argv[2]) if len(sys.argv) > 2 else 180)
shutil.move(os.path.join(out, os.path.basename(path)), path)
formulas = sum(1 for ws in load_workbook(path).worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
errs = [f"{ws.title}!{c.coordinate}" for ws in load_workbook(path, data_only=True).worksheets for row in ws.iter_rows() for c in row if c.value in ERRORS]
print(json.dumps({"status": "success" if not errs else "errors_found", "total_formulas": formulas, "total_errors": len(errs), "errors": errs[:50]}))
